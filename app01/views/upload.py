import os
from openpyxl import load_workbook
from datetime import datetime

from django.conf import settings
from django.shortcuts import render, redirect
from django import forms
from django.db import models

from app01.models import UserInfo, Department, User
from app01.forms.bootstrap_forms import BootstrapForm, BootstrapModelForm

import logging

# 创建日志对象
logger = logging.getLogger()


def upload_file(request):
    if request.method == "GET":
        return render(request, "file/upload_add.html", {'status': False})
    # upload_image(request)
    upload_excel(request)
    return render(request, "file/upload_add.html", {'status': True})


def upload_image(request):
    save_file_to_file(request)


def upload_excel(request):
    # 方式一：将上传的文件内容保存到文件
    # save_file_to_file(request)
    # 方式二：将上传的Excel文件内容保存到数据库中
    save_excel_to_db(request)


def save_file_to_file(request):
    files = request.FILES
    # 获取静态文件目录（在settings.py文件中 STATICFILES_DIRS 配置为[BASE_DIR / 'app01']）
    static_dir = settings.STATICFILES_DIRS[0]
    # 获取存放文件的绝对路径，例如 app01/upload/
    upload_dir = os.path.join(static_dir, 'upload')
    os.makedirs(upload_dir, exist_ok=True)  # 自动创建目录
    file_obj = files.get("file")
    # 获取上传的文件名
    file_name = file_obj.name
    # file_name = "upload_file01.png"
    file_path = os.path.join(upload_dir, file_name)
    with open(file=file_path, mode="wb") as f:
        for chunk in file_obj.chunks():
            f.write(chunk)


def save_excel_to_db(request):
    files = request.FILES
    file_obj = files.get("file")
    wb = load_workbook(file_obj)
    # 获取Excel中的第一个sheet页
    sheet = wb.worksheets[0]
    # 表头验证
    expected_headers = ['用户名', '密码', '年龄', '账户余额', '入职时间', '性别', '所属部门']
    actual_headers = [cell.value for cell in sheet[1]]
    if actual_headers != expected_headers:
        raise ValueError("Excel表头结构不匹配")
    # 数据批处理
    batch_size = 1000
    instances = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        try:
            is_exists = UserInfo.objects.filter(username=row[0]).exists()
            if is_exists:
                # print(f"数据({row[0]})已存在")
                logger.info(f"数据({row[0]})已存在")
                continue
            depart_id = int(row[6])
            if not Department.objects.filter(id=depart_id).exists():
                raise ValueError(f"部门ID({depart_id})不存在")
            instances.append(UserInfo(
                username=row[0],
                password=row[1],  # 实际生产需加密处理
                age=int(row[2]),
                account=float(str(row[3]).strip()) if row[3] else 0.0,
                create_time=row[4].date() if isinstance(row[4], datetime) else datetime.strptime(row[4],
                                                                                                 '%Y-%m-%d').date(),
                gender=1 if row[5] == '男' else 2,
                depart_id=int(row[6])
            ))
            if len(instances) >= batch_size:
                UserInfo.objects.bulk_create(instances)
                instances = []
        except Exception as e:
            # print(f"{row[0]}-此行数据异常: {str(e)}")
            logger.error(f"{row[0]}-此行数据异常: {str(e)}")
            continue
    if instances:
        UserInfo.objects.bulk_create(instances)


class UserForm(BootstrapForm):
    exclude_fields = ['avatar']

    name = forms.CharField(label='姓名')
    age = forms.IntegerField(label='年龄')
    avatar = forms.FileField(label='头像')


class UserModelForm(BootstrapModelForm):
    exclude_fields = ['avatar']

    class Meta:
        model = User
        fields = "__all__"


def upload_add(request):
    # return upload_user_by_form(request)
    return upload_user_by_model_form(request)


def upload_user_by_form(request):
    form = UserForm()
    if request.method == "GET":
        return render(request, "file/upload_add.html", {'status': False, 'form': form})
    try:
        form = UserForm(data=request.POST, files=request.FILES)
        if form.is_valid():
            name = form.cleaned_data['name']
            age = form.cleaned_data['age']
            # 获取上传的头像文件对象
            file_obj = form.cleaned_data['avatar']
            # 保存在数据库中的文件路径
            # db_file_path = os.path.join('upload', 'avatar', file_obj.name)
            # 实际存放头像图片的文件路径
            # file_path = os.path.join(settings.STATICFILES_DIRS[0], db_file_path)
            file_path = os.path.join("media", file_obj.name)
            if not os.path.exists(file_path):
                os.makedirs(os.path.dirname(file_path), exist_ok=True)
            with open(file=file_path, mode="wb") as f:
                for chunk in file_obj.chunks():
                    f.write(chunk)
            User.objects.create(name=name, age=age, avatar=file_path)
            logger.info("上传用户信息成功")
            return redirect("/upload/list/")
        return render(request, "file/upload_add.html", {'status': False, 'form': form})
    except Exception as e:
        logger.error(f"上传用户信息失败：{e}")
        return render(request, "file/upload_add.html", {'status': False, 'form': form})


def upload_user_by_model_form(request):
    form = UserModelForm()
    if request.method == "GET":
        return render(request, "file/upload_add.html", {'status': False, 'form': form})
    try:
        form = UserModelForm(data=request.POST, files=request.FILES)
        if form.is_valid():
            # 会自动将图片保存到media目录下，数据库中保存的是不带media前缀的图片路径，例如：avatar/三国演义.jpg
            form.save()
            return redirect("/upload/list/")
        return render(request, "file/upload_add.html", {'status': False, 'form': form})
    except Exception as e:
        logger.error(f"上传失败：{e}")
        return render(request, "file/upload_add.html", {'status': False, 'form': form})


def upload_list(request):
    queryset = User.objects.all()
    return render(request, "file/upload_list.html", {'queryset': queryset})


def upload_edit(request, id):
    data_obj = User.objects.get(id=id)
    if not data_obj:
        logger.warning(f"数据ID({id})不存在")
        return redirect("/upload/list/")
    if request.method == "GET":
        # 将查询到的数据封装到ModelForm中，用于在界面上展示修改之前的数据
        form = UserModelForm(instance=data_obj)
        return render(request, "file/upload_edit.html", {'form': form})
    form = UserModelForm(data=request.POST, files=request.FILES, instance=data_obj)
    if form.is_valid():
        form.save()
        return redirect("/upload/list/")
    return render(request, "file/upload_edit.html", {'form': form})


def upload_delete(request, id):
    is_exists = User.objects.filter(id=id).exists()
    if not is_exists:
        logger.error(f"数据ID({id})不存在")
        raise RuntimeError(f"数据ID({id})不存在")
    User.objects.get(id=id).delete()
    return redirect("/upload/list/")
